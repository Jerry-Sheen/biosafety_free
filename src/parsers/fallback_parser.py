"""
Lightweight fallback document parser.
Provides rapid, dependable parsing for PDF, DOCX, XLSX, CSV, and text files.
Also handles Korean character encodings (UTF-8, CP949, EUC-KR) automatically.
"""
from pathlib import Path
from typing import Dict, Any, List
from src.security.logger import get_secure_logger

logger = get_secure_logger("fallback_parser")


class FallbackParser:
    """Rapid parser for common document formats."""

    @classmethod
    def parse(cls, file_path: Path) -> Dict[str, Any]:
        file_path = Path(file_path)
        ext = file_path.suffix.lower()

        if ext == ".pdf":
            return cls._parse_pdf(file_path)
        elif ext in (".docx", ".doc"):
            return cls._parse_docx(file_path)
        elif ext in (".xlsx", ".xls"):
            return cls._parse_excel(file_path)
        elif ext == ".csv":
            return cls._parse_csv(file_path)
        elif ext in (".txt", ".md", ".json"):
            return cls._parse_text(file_path)
        else:
            raise ValueError(f"Unsupported file format for fallback parser: {ext}")

    @classmethod
    def _parse_pdf(cls, file_path: Path) -> Dict[str, Any]:
        import pypdf
        logger.info(f"Fallback parsing PDF: {file_path.name}")
        reader = pypdf.PdfReader(str(file_path))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"--- Page {i + 1} ---\n{text}")

        combined_text = "\n\n".join(pages_text)
        is_empty = len(combined_text.strip()) == 0

        metadata = {
            "source": file_path.name,
            "format": "pdf",
            "page_count": len(reader.pages),
            "parser": "fallback_pypdf",
            "is_scanned_image": is_empty
        }

        if is_empty:
            warning_msg = (
                f"[주의: {file_path.name} 파일에 텍스트가 없습니다. "
                "스캔된 이미지 PDF이므로 Docling OCR 모드를 활성화하여 처리해야 합니다.]"
            )
            combined_text = warning_msg

        return {
            "text": combined_text,
            "metadata": metadata,
            "tables": []
        }

    @classmethod
    def _parse_docx(cls, file_path: Path) -> Dict[str, Any]:
        import docx
        logger.info(f"Fallback parsing DOCX: {file_path.name}")
        doc = docx.Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Extract docx tables
        tables_md = []
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
            if rows:
                header = "| " + " | ".join(rows[0]) + " |"
                sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
                data_rows = ["| " + " | ".join(r) + " |" for r in rows[1:]]
                tables_md.append("\n".join([header, sep] + data_rows))

        full_text = "\n\n".join(paragraphs)
        if tables_md:
            full_text += "\n\n### Tables in Document:\n" + "\n\n".join(tables_md)

        return {
            "text": full_text,
            "metadata": {
                "source": file_path.name,
                "format": "docx",
                "paragraph_count": len(paragraphs),
                "table_count": len(tables_md),
                "parser": "fallback_docx"
            },
            "tables": tables_md
        }

    @classmethod
    def _parse_excel(cls, file_path: Path) -> Dict[str, Any]:
        import pandas as pd
        logger.info(f"Fallback parsing Excel: {file_path.name}")
        excel_file = pd.ExcelFile(str(file_path))
        sheet_texts = []
        tables_md = []

        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            df = df.dropna(how="all")
            markdown_table = df.to_markdown(index=False)
            sheet_content = f"### Sheet: {sheet_name}\n\n{markdown_table}"
            sheet_texts.append(sheet_content)
            tables_md.append(markdown_table)

        return {
            "text": "\n\n".join(sheet_texts),
            "metadata": {
                "source": file_path.name,
                "format": "excel",
                "sheets": excel_file.sheet_names,
                "parser": "fallback_excel"
            },
            "tables": tables_md
        }

    @classmethod
    def _parse_csv(cls, file_path: Path) -> Dict[str, Any]:
        import pandas as pd
        logger.info(f"Fallback parsing CSV: {file_path.name}")

        # Attempt UTF-8, CP949, EUC-KR for Korean CSV support
        df = None
        for enc in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
            try:
                df = pd.read_csv(str(file_path), encoding=enc)
                logger.info(f"Successfully read CSV using encoding: {enc}")
                break
            except Exception:
                continue

        if df is None:
            df = pd.read_csv(str(file_path), encoding_errors="replace")

        markdown_table = df.to_markdown(index=False)
        return {
            "text": f"### CSV Data: {file_path.name}\n\n{markdown_table}",
            "metadata": {
                "source": file_path.name,
                "format": "csv",
                "row_count": len(df),
                "column_count": len(df.columns),
                "parser": "fallback_csv"
            },
            "tables": [markdown_table]
        }

    @classmethod
    def _parse_text(cls, file_path: Path) -> Dict[str, Any]:
        logger.info(f"Fallback parsing text file: {file_path.name}")
        content = None
        for enc in ("utf-8", "cp949", "euc-kr", "latin1"):
            try:
                content = file_path.read_text(encoding=enc)
                break
            except Exception:
                continue

        if content is None:
            content = file_path.read_text(encoding="utf-8", errors="replace")

        return {
            "text": content,
            "metadata": {
                "source": file_path.name,
                "format": file_path.suffix.lstrip("."),
                "char_count": len(content),
                "parser": "fallback_text"
            },
            "tables": []
        }
